"""PPTX handler: anonymizes text in shapes (preserves layout)."""
import os
import sys
from pptx import Presentation

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from presidio_helper import anonymize_text


def anonymize(path, output=None, template=None):
    prs = Presentation(path)
    all_texts = []

    for slide in prs.slides:
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            for para in shape.text_frame.paragraphs:
                # Join runs to anonymize coherent text, then rewrite first run
                full_text = ''.join(run.text for run in para.runs)
                if not full_text.strip():
                    continue
                anon = anonymize_text(full_text)
                all_texts.append(anon)
                if anon != full_text and para.runs:
                    para.runs[0].text = anon
                    for run in para.runs[1:]:
                        run.text = ''

    if output:
        prs.save(output)
        return output
    return '\n'.join(all_texts)
