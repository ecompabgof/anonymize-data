"""XLSX handler: removes sensitive columns + anonymizes free-text cells.

Strategy:
  1. If template provided: use column list from YAML.
  2. Else: auto-detect via header keyword match.
  3. Remove sensitive columns entirely.
  4. Anonymize remaining text cells (>10 chars) via Presidio.
"""
import os
import sys
from pathlib import Path
import openpyxl
import yaml

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from presidio_helper import anonymize_text

SENSITIVE_KEYWORDS = [
    'nom', 'prenom', 'prénom', 'name', 'firstname', 'lastname', 'fullname',
    'adresse', 'address', 'addr',
    'email', 'mail', 'e-mail',
    'telephone', 'téléphone', 'tel', 'phone', 'mobile', 'portable',
    'iban', 'rib', 'bic', 'swift',
    'nir', 'securite sociale', 'sécurité sociale', 'ssn', 'numero ss',
    'date naissance', 'birth', 'birthday', 'naissance',
    'passport', 'passeport', 'cni', 'carte identite',
    'cb', 'card', 'carte bancaire', 'cvv',
]


def _auto_detect_cols(ws):
    headers = [(c.value or '').strip().lower() for c in ws[1]]
    cols = []
    for i, h in enumerate(headers, start=1):
        if not h:
            continue
        if any(kw in h for kw in SENSITIVE_KEYWORDS):
            cols.append(i)
    return cols


def _load_template(template_name):
    tpl_path = Path(__file__).parent.parent.parent / 'templates' / f'{template_name}.yml'
    if not tpl_path.exists():
        return None
    return yaml.safe_load(tpl_path.read_text(encoding='utf-8'))


def _cols_from_template(ws, tpl):
    sensitive_names = [c.lower().strip() for c in tpl.get('sensitive_columns', [])]
    headers = [(c.value or '').strip().lower() for c in ws[1]]
    return [i + 1 for i, h in enumerate(headers) if h and h in sensitive_names]


def anonymize(path, output=None, template=None):
    wb = openpyxl.load_workbook(path)
    tpl = _load_template(template) if template else None

    for ws in wb.worksheets:
        cols = _cols_from_template(ws, tpl) if tpl else _auto_detect_cols(ws)

        # Delete sensitive columns (reverse to preserve indices)
        for col in sorted(cols, reverse=True):
            ws.delete_cols(col)

        # Anonymize remaining text cells (long strings only — skip numbers/short codes)
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, str) and len(cell.value) > 10:
                    cell.value = anonymize_text(cell.value)

    if output:
        wb.save(output)
        return output

    # Inline mode: dump as text grid
    out_lines = []
    for ws in wb.worksheets:
        out_lines.append(f'=== Sheet: {ws.title} ===')
        for row in ws.iter_rows(values_only=True):
            out_lines.append(' | '.join(str(c) if c is not None else '' for c in row))
        out_lines.append('')
    return '\n'.join(out_lines)
